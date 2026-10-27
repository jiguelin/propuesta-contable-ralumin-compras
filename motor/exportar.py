# -*- coding: utf-8 -*-
"""
motor/exportar.py — Excel de cada bloque para Contasis, su Excel de revisión y el resumen del mes.

Excel del bloque: plantilla NewContaSis de 50 columnas (A–AX), solo filas de datos, en el orden del archivador
(del más reciente al más antiguo); al final las facturas con comisión o retención. Diferencias con el app de compras:
  · AH (cuenta total) = la cuenta de pago de CADA fila (1011, 104103, 104102, 4212…).
  · AF = cuenta de gasto elegida; AE = condición CON/CRE de la fila; AS = glosa con el estilo de Ralumin.
  · AK–AP no llevan la retención (se completa con el reporte AJUSTES DE PAGO); detracción y percepción, igual que hoy.
"""
from __future__ import annotations

import io
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ralumin import config as cfg

from . import pcge
from .excel_newcontasis import CABECERAS, _escribir_filas, fila_newcontasis
from .sesion import Bloque, Fila, Sesion

AZUL = '1F3B5C'
COLORES = {'ok': 'C6EFCE', 'amarillo': 'FFEB9C', 'rojo': 'FFC7CE'}
COL_AE, COL_AF, COL_AS, COL_AV = 30, 31, 44, 47
COL_AK, COL_AP = 36, 41


def fila_contasis(f: Fila, plan=None) -> list:
    p, d = f.p, f.d
    p.cuenta = d.cuenta_gasto
    p.glosa = d.glosa
    p.av = pcge.av_por_cuenta(d.cuenta_gasto)
    fila = fila_newcontasis(p, d.cuenta_pago, usd_en_soles=True, cuenta_percepcion='40113')
    fila[COL_AE] = d.condicion or fila[COL_AE]
    fila[COL_AF] = d.cuenta_gasto
    fila[COL_AS] = d.glosa[:60]
    fila[COL_AV] = p.av
    if fila[COL_AK] == 3:                     # retención declarada en el XML: va por el reporte de ajustes
        for j in range(COL_AK, COL_AP + 1):
            fila[j] = ''
    return fila


def excel_bloque(s: Sesion, b: Bloque) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = 'COMPRAS'
    _escribir_filas(ws, [fila_contasis(f) for f in s.filas_bloque(b)], 1)
    for j in range(1, 51):
        ws.column_dimensions[get_column_letter(j)].width = 14
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------------------------- helpers de hojas
def _cab(ws, fila: list, fila_n: int | None = None):
    ws.append(fila)
    for cell in ws[ws.max_row if fila_n is None else fila_n]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=AZUL)
        cell.alignment = Alignment(wrap_text=True, vertical='center')


def _titulo(ws, texto: str, sub: str = ''):
    ws.append([texto])
    ws['A1'].font = Font(bold=True, size=13)
    if sub:
        ws.append([sub])
        ws['A2'].font = Font(italic=True, color='555555')
    ws.append([])


def _anchos(ws, anchos: list[int]):
    for j, w in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w


def _montos(ws, desde_fila: int, columnas: list[int]):
    for row in ws.iter_rows(min_row=desde_fila):
        for j in columnas:
            if j - 1 < len(row) and isinstance(row[j - 1].value, (int, float)):
                row[j - 1].number_format = '#,##0.00'


def _desc(plan, cuenta: str) -> str:
    cu = plan.get(cuenta) if plan and cuenta else None
    return cu.descripcion if cu else ''


def _asiento(b: Bloque, k: int) -> str:
    return str(b.asiento_inicial + k) if b.asiento_inicial else ''


# ----------------------------------------------------------------------------- revisión del bloque
def excel_revision_bloque(s: Sesion, b: Bloque, plan=None) -> bytes:
    filas = s.filas_bloque(b)
    wb = Workbook()

    # ---- PROPUESTA ----
    ws = wb.active
    ws.title = 'PROPUESTA'
    _titulo(ws, f'{cfg.EMPRESA} — {b.codigo} {b.nombre} — {s.periodo}', f'{len(filas)} comprobantes · creado {b.creado.replace("T", " ")}')
    _cab(ws, ['#', 'Asiento', 'Grupo', 'Fecha', 'Tipo', 'Serie-Número', 'RUC', 'Proveedor', 'Moneda', 'Base', 'IGV', 'Total', 'TC', 'Total S/',
              'Cuenta gasto', 'Descripción', 'Reparto 60/40', 'Cuenta pago', 'Descripción pago', 'Cond.', 'Retención S/', 'Comisión', 'Glosa', 'Avisos'])
    for k, f in enumerate(filas):
        c, d = f.c, f.d
        ws.append([k + 1, _asiento(b, k), cfg.GRUPO_NOMBRE.get(d.grupo, d.grupo), c.fecha_emision, c.tipo_codigo, c.serie_numero, c.ruc_emisor,
                   c.nombre_emisor, c.moneda_codigo, float(c.gravadas or c.valor_venta), float(c.igv), f.total, f.tc or '', f.total_soles,
                   d.cuenta_gasto, _desc(plan, d.cuenta_gasto), 'SÍ' if d.reparto else '', d.cuenta_pago, _desc(plan, d.cuenta_pago), d.condicion,
                   d.retencion_monto if d.retencion else '', d.comision or '', d.glosa, ' | '.join(f.avisos)])
        ws.cell(ws.max_row, 1).fill = PatternFill('solid', fgColor=COLORES[f.nivel])
    _anchos(ws, [5, 8, 18, 11, 5, 16, 13, 34, 7, 11, 10, 11, 7, 11, 11, 30, 8, 10, 26, 6, 11, 9, 40, 70])
    _montos(ws, 5, [10, 11, 12, 14, 21, 22])
    ws.freeze_panes = 'A5'

    # ---- AJUSTES DE PAGO ----
    aj = [(k, f) for k, f in enumerate(filas) if f.tiene_ajuste]
    ws = wb.create_sheet('AJUSTES DE PAGO')
    _titulo(ws, 'Completar en Contasis después de importar este bloque',
            f'Por cada fila: abrir el asiento, agregar {cfg.CUENTA_COMISION} al Debe (comisión) y {cfg.CUENTA_RETENCION} al Haber (retención), '
            'y cambiar la cuenta 10 al monto final.')
    _cab(ws, ['#', 'Asiento', 'Comprobante', 'Proveedor', 'Cuenta pago', 'Moneda', 'Cuenta 10 al importar S/',
              f'Agregar Debe {cfg.CUENTA_COMISION} S/', f'Agregar Haber {cfg.CUENTA_RETENCION} S/', 'Cuenta 10 final S/',
              'Comisión US$', 'Cuenta 10 final US$', 'Comprobante de retención', 'Fecha retención'])
    tot_com = tot_ret = 0.0
    for k, f in aj:
        d = f.d
        usd = f.c.moneda_codigo == 'USD'
        tc = f.tc or 1.0
        com_s = round(d.comision * (tc if usd else 1.0), 2)
        ret = round(d.retencion_monto if d.retencion else 0.0, 2)
        final_s = round(f.total_soles - ret + com_s, 2)
        tot_com += com_s
        tot_ret += ret
        ws.append([k + 1, _asiento(b, k), f.c.serie_numero, f.c.nombre_emisor, d.cuenta_pago, f.c.moneda_codigo, f.total_soles,
                   com_s or '', ret or '', final_s, d.comision if usd and d.comision else '',
                   round(f.total - ret / tc + d.comision, 2) if usd else '',
                   f'{d.retencion_serie}-{d.retencion_numero}' if d.retencion_serie else ('FALTA' if d.retencion else ''),
                   d.retencion_fecha])
    ws.append(['', '', '', 'TOTAL', '', '', '', round(tot_com, 2), round(tot_ret, 2)])
    ws.cell(ws.max_row, 4).font = Font(bold=True)
    _anchos(ws, [5, 8, 16, 34, 11, 7, 14, 14, 14, 14, 11, 13, 18, 12])
    _montos(ws, 5, [7, 8, 9, 10, 11, 12])
    ws.freeze_panes = 'A5'

    # ---- REPARTO 60-40 ----
    rep = [(k, f) for k, f in enumerate(filas) if f.d.reparto]
    ws = wb.create_sheet('REPARTO 60-40')
    _titulo(ws, f'Gastos que Ralumin reparte {cfg.PORCENTAJE_ADM:.0f} % ADM / {100 - cfg.PORCENTAJE_ADM:.0f} % VTAS',
            'Se importaron al 100 % en la cuenta ADM. En Contasis, partir la base así:')
    _cab(ws, ['#', 'Asiento', 'Comprobante', 'Proveedor', 'Base S/', 'Cuenta ADM', 'Monto ADM S/', 'Cuenta VTAS', 'Monto VTAS S/'])
    for k, f in rep:
        base = round(float(f.c.gravadas or f.c.valor_venta) * (f.tc if f.c.moneda_codigo == 'USD' and f.tc else 1.0), 2)
        adm = round(base * cfg.PORCENTAJE_ADM / 100, 2)
        ws.append([k + 1, _asiento(b, k), f.c.serie_numero, f.c.nombre_emisor, base, f.d.cuenta_gasto, adm, cfg.par_vtas(f.d.cuenta_gasto), round(base - adm, 2)])
    _anchos(ws, [5, 8, 16, 34, 11, 11, 12, 11, 12])
    _montos(ws, 5, [5, 7, 9])

    # ---- PAGOS AGRUPADOS ----
    grupos = defaultdict(list)
    for k, f in enumerate(filas):
        if f.d.pago_agrupado:
            grupos[(f.c.ruc_emisor, f.d.fecha_pago, f.d.cuenta_pago)].append(f)
    ws = wb.create_sheet('PAGOS AGRUPADOS')
    _titulo(ws, 'Pagos que cubren varias facturas o facturas pagadas en partes', 'Para cuadrar con el extracto del banco.')
    _cab(ws, ['RUC', 'Proveedor', 'Fecha de pago', 'Cuenta', 'Comprobantes', 'Suma de facturas S/'])
    for (ruc, fp, cta), fs in grupos.items():
        ws.append([ruc, fs[0].c.nombre_emisor, fp, cta, ', '.join(x.c.serie_numero for x in fs), round(sum(x.total_soles for x in fs), 2)])
    _anchos(ws, [13, 34, 12, 10, 40, 16])
    _montos(ws, 5, [6])

    # ---- PROVEEDORES NUEVOS ----
    ws = wb.create_sheet('PROVEEDORES NUEVOS')
    _titulo(ws, 'Proveedores que no están en Entidades de Contasis', 'Crearlos antes de importar el bloque.')
    _cab(ws, ['RUC', 'Razón social', 'Dirección'])
    vistos = set()
    for f in filas:
        if f.proveedor_nuevo and f.c.ruc_emisor not in vistos:
            vistos.add(f.c.ruc_emisor)
            ws.append([f.c.ruc_emisor, f.c.nombre_emisor, f.c.direccion_emisor])
    _anchos(ws, [13, 44, 60])

    # ---- ORDEN DE SELLOS ----
    ws = wb.create_sheet('ORDEN DE SELLOS')
    _titulo(ws, 'Orden de los comprobantes en el Excel (y asiento esperado si se indicó el primero)')
    _cab(ws, ['#', 'Asiento', 'Fecha', 'Comprobante', 'Proveedor', 'Total', 'Moneda'])
    for k, f in enumerate(filas):
        ws.append([k + 1, _asiento(b, k), f.c.fecha_emision, f.c.serie_numero, f.c.nombre_emisor, f.total, f.c.moneda_codigo])
    _anchos(ws, [5, 8, 11, 16, 40, 11, 7])
    _montos(ws, 4, [6])

    # ---- CONTASIS (ref) ----
    ws = wb.create_sheet('CONTASIS (ref)')
    ws.append(CABECERAS)
    for cell in ws[1]:
        cell.font = Font(bold=True, size=8)
        cell.alignment = Alignment(wrap_text=True)
    _escribir_filas(ws, [fila_contasis(f) for f in filas], 2)
    ws.freeze_panes = 'A2'

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------------------------- resumen del mes
def excel_mes(s: Sesion, plan=None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = 'AVANCE'
    av = s.avance()
    _titulo(ws, f'{cfg.EMPRESA} — Compras {s.periodo}', f'{av["contabilizado"]} de {av["total"]} comprobantes contabilizados')
    _cab(ws, ['Grupo', 'Pendiente', 'Exportado', 'Contabilizado', 'Pendiente S/', 'Pendiente US$'])
    for g in cfg.GRUPOS:
        fs = [f for f in s.filas.values() if f.d.grupo == g.id and f.d.estado != 'excluido']
        if not fs:
            continue
        cnt = {e: sum(1 for f in fs if f.d.estado == e) for e in ('pendiente', 'exportado', 'contabilizado')}
        pen = [f for f in fs if f.d.estado == 'pendiente']
        ws.append([g.nombre, cnt['pendiente'], cnt['exportado'], cnt['contabilizado'],
                   round(sum(f.total for f in pen if f.c.moneda_codigo != 'USD'), 2), round(sum(f.total for f in pen if f.c.moneda_codigo == 'USD'), 2)])
    _anchos(ws, [24, 11, 11, 13, 14, 14])
    _montos(ws, 5, [5, 6])

    ws = wb.create_sheet('BLOQUES')
    _cab(ws, ['Bloque', 'Nombre', 'Estado', 'Comprobantes', 'Total S/', 'Total US$', 'Creado', 'Confirmado', 'Asiento inicial'])
    for b in s.bloques:
        fs = s.filas_bloque(b)
        ws.append([b.codigo, b.nombre, b.estado, len(fs), round(sum(f.total for f in fs if f.c.moneda_codigo != 'USD'), 2),
                   round(sum(f.total for f in fs if f.c.moneda_codigo == 'USD'), 2), b.creado.replace('T', ' '), b.confirmado.replace('T', ' '),
                   b.asiento_inicial or ''])
    _anchos(ws, [8, 34, 13, 13, 12, 12, 18, 18, 13])
    _montos(ws, 2, [5, 6])

    for estado, titulo in (('pendiente', 'PENDIENTES'), ('excluido', 'EXCLUIDOS')):
        ws = wb.create_sheet(titulo)
        _cab(ws, ['Grupo', 'Fecha', 'Serie-Número', 'RUC', 'Proveedor', 'Moneda', 'Total', 'Cuenta gasto', 'Cuenta pago', 'Motivo / avisos', 'Archivo'])
        for f in s.lista(estado):
            ws.append([cfg.GRUPO_NOMBRE.get(f.d.grupo, ''), f.c.fecha_emision, f.c.serie_numero, f.c.ruc_emisor, f.c.nombre_emisor, f.c.moneda_codigo,
                       f.total, f.d.cuenta_gasto, f.d.cuenta_pago, f.d.motivo or ' | '.join(f.avisos), f.c.archivo])
        _anchos(ws, [20, 11, 16, 13, 34, 7, 11, 11, 10, 70, 30])
        _montos(ws, 2, [7])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
