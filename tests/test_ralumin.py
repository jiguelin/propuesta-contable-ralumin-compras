# -*- coding: utf-8 -*-
"""Pruebas del motor de Compras Ralumin con datos ficticios.   python -m pytest -q"""
from __future__ import annotations

import io
import json
from datetime import date

import openpyxl
import pytest

from motor.almacen import Almacen
from motor.diario import construir_memoria, leer_diario, normalizar_serie_numero
from motor.exportar import excel_bloque, excel_mes, excel_revision_bloque
from motor.maestros import leer_entidades
from motor.plan_cuentas import leer_plan_excel
from motor.retenciones import leer_padron, leer_xml_retencion
from motor.sesion import Maestros, Sesion
from motor.tipo_cambio import TC, HistorialTC
from tests.fabrica import diario_excel, plan_excel, xml_factura, xml_retencion

CUENTAS = [('1011', 'CAJA MN'), ('1012', 'CAJA ME'), ('104102', 'CUENTA CORRIENTE ME-191'), ('104103', '191 - BCP - MN'),
           ('104105', 'BCO BBVA -MN'), ('4212', 'EMITIDAS - Facturas por pagar'), ('401111', 'IGV'), ('401141', 'IGV-Retenciones por Pagar'),
           ('6011020', 'MERCADERIAS - Compras'), ('6031025', 'MATERIALES AUXILIARES - Compras'), ('6315094', 'OTROS GASTOS DE VIAJE - ADM'),
           ('6315095', 'OTROS GASTOS DE VIAJE - VTAS'), ('6560094', 'SUMINISTROS - ADM'), ('6560095', 'SUMINISTROS - VTAS'),
           ('656009401', 'GASOLINA - ADM'), ('656009501', 'GASOLINA - VTAS'), ('6391094', 'GASTOS BANCARIOS - ADM'),
           ('609192012', 'GASTOS DE NACIONALIZACION'), ('6399094', 'OTROS SERVICIOS PRESTADOS POR TERCEROS - ADM')]

PROV_MAT, PROV_CAJA, PROV_USD, PROV_PEAJE = '20999999901', '20999999902', '20999999903', '20523621212'


def _tc(*dias):
    return HistorialTC([TC(date(2026, 7, d), 3.39, 3.40) for d in range(1, 32)])


def maestros(con_diario=True, padron=None, retenciones=None) -> Maestros:
    plan = leer_plan_excel(plan_excel(CUENTAS))
    mem = None
    if con_diario:
        asientos = [
            # proveedor de materiales pagado por BCP en junio
            {'mes': '06', 'n': 10, 'ruc': PROV_MAT, 'doc': '01', 'numero': '00E001-0000000000100', 'fecha': date(2026, 6, 10),
             'lineas': [('6031025', 1000, 0), ('401111', 180, 0), ('4212', 0, 1180), ('4212', 1180, 0), ('104103', 0, 1180)]},
            # proveedor de caja chica pagado en efectivo
            {'mes': '06', 'n': 11, 'ruc': PROV_CAJA, 'doc': '01', 'numero': '00F010-0000000000050', 'fecha': date(2026, 6, 11),
             'lineas': [('6560094', 60, 0), ('6560095', 40, 0), ('401111', 18, 0), ('1011', 0, 118)]},
            # una factura de julio que YA está contabilizada
            {'mes': '07', 'n': 1, 'ruc': PROV_MAT, 'doc': '01', 'numero': '00E001-0000000000150', 'fecha': date(2026, 7, 2),
             'lineas': [('6031025', 100, 0), ('401111', 18, 0), ('104103', 0, 118)]},
        ]
        mem = construir_memoria(leer_diario(diario_excel(asientos)))
    return Maestros(plan=plan, tc=_tc(), memoria=mem, entidades={PROV_MAT: 'MATERIALES SAC', PROV_CAJA: 'FERRE SAC', PROV_PEAJE: 'LIMA EXPRESA'},
                    padron=padron or set(), retenciones=retenciones or [])


def sesion(*xmls, m=None) -> tuple[Sesion, Maestros]:
    s = Sesion('2026-07')
    s.agregar_xml([(f'x{k}.xml', x) for k, x in enumerate(xmls)])
    m = m or maestros()
    s.procesar(m)
    return s, m


def fila(s, serie):
    return next(f for f in s.filas.values() if f.c.serie_numero == serie)


# ----------------------------------------------------------------------------- diario
def test_diario_memoria_y_ya_contabilizado():
    m = maestros()
    assert m.memoria.pago_reciente(PROV_MAT) == '104103'
    assert m.memoria.pago_reciente(PROV_CAJA) == '1011'
    assert m.memoria.suele_repartir(PROV_CAJA) is True
    assert normalizar_serie_numero('00E001-0000000000150') == ('E001', '150')
    s, _ = sesion(xml_factura(PROV_MAT, 'MATERIALES SAC', 'E001-150', '2026-07-02', 100))
    assert fila(s, 'E001-150').d.estado == 'contabilizado'          # ya estaba en el Diario: no vuelve a la propuesta


# ----------------------------------------------------------------------------- grupos y cuentas
def test_grupos_y_cuentas_de_pago():
    s, _ = sesion(
        xml_factura(PROV_MAT, 'MATERIALES SAC', 'E001-185', '2026-07-31', 6440.69, 1159.32, items=[('FIERRO 3/8', 6440.69)]),
        xml_factura(PROV_CAJA, 'FERRE SAC', 'F010-60', '2026-07-23', 593.22, 106.78),
        xml_factura(PROV_USD, 'PARTS EIRL', 'E001-53', '2026-07-01', 1000, moneda='USD', forma_pago='Credito', cuotas=[('2026-09-29', 1180)]),
        xml_factura(PROV_USD, 'PARTS EIRL', 'E001-54', '2026-07-01', 1000, moneda='USD', forma_pago='Credito', cuotas=[('2026-07-31', 1180)]),
        xml_factura(PROV_USD, 'PARTS EIRL', 'E001-55', '2026-07-30', 300, moneda='USD'),
        xml_factura(PROV_PEAJE, 'LIMA EXPRESA S.A.C.', 'F805-1873721', '2026-07-31', 5.34, 0.96, items=[('PEAJE VIA 105', 5.34)]),
    )
    assert fila(s, 'E001-185').d.grupo == 'soles_contado' and fila(s, 'E001-185').d.cuenta_pago == '104103'
    assert fila(s, 'F010-60').d.grupo == 'caja_chica' and fila(s, 'F010-60').d.cuenta_pago == '1011'
    assert fila(s, 'E001-53').d.grupo == 'usd_60_120' and fila(s, 'E001-53').d.cuenta_pago == '4212'
    assert fila(s, 'E001-54').d.grupo == 'usd_credito'
    assert fila(s, 'E001-55').d.grupo == 'usd_contado' and fila(s, 'E001-55').d.cuenta_pago == '104102'
    peaje = fila(s, 'F805-1873721')
    assert peaje.d.cuenta_gasto == '6315095' and not peaje.d.reparto            # peajes: una sola cuenta
    assert peaje.d.glosa.startswith('GASTOS DE VIAJE - PEAJE')
    assert fila(s, 'E001-185').d.cuenta_gasto == '6031025'                      # memoria del proveedor


def test_contado_nunca_4212():
    s, m = sesion(xml_factura(PROV_CAJA, 'FERRE SAC', 'F010-61', '2026-07-20', 100))
    f = fila(s, 'F010-61')
    s.editar(f.id, m, cuenta_pago='4212')
    assert f.nivel == 'rojo' and any('contado con 4212' in a.lower() for a in f.avisos)


def test_reparto_por_historial():
    s, _ = sesion(xml_factura(PROV_CAJA, 'FERRE SAC', 'F010-62', '2026-07-20', 50, items=[('UTILES DE OFICINA', 50)]))
    f = fila(s, 'F010-62')
    assert f.d.cuenta_gasto == '6560094' and f.d.reparto


# ----------------------------------------------------------------------------- retenciones
def test_retencion_regla_700_y_exclusiones():
    s, m = sesion(
        xml_factura(PROV_MAT, 'MATERIALES SAC', 'E001-185', '2026-07-31', 6440.69, 1159.32),                  # 7,600.01 → retiene 228.00
        xml_factura(PROV_CAJA, 'FERRE SAC', 'F010-70', '2026-07-15', 400),                                     # 472 + 472 mismo día → 944 > 700
        xml_factura(PROV_CAJA, 'FERRE SAC', 'F010-71', '2026-07-15', 400),
        xml_factura('20999999904', 'BUEN SAC', 'F001-1', '2026-07-15', 2000,
                    leyendas=[('', 'NO RETENER, BUEN CONTRIBUYENTE RESOLUCION 0230050222614')]),
        xml_factura('20999999905', 'CHICO SAC', 'F001-2', '2026-07-15', 500),                                  # 590 ≤ 700
    )
    f = fila(s, 'E001-185')
    assert f.retencion_sugerida and f.d.retencion and f.d.retencion_monto == 228.00
    assert fila(s, 'F010-70').d.retencion and fila(s, 'F010-71').d.retencion
    assert not fila(s, 'F001-1').d.retencion and 'No Retener' in fila(s, 'F001-1').retencion_exclusion
    assert not fila(s, 'F001-2').d.retencion
    # padrón SUNAT: excluye
    s2, _ = sesion(xml_factura(PROV_MAT, 'MATERIALES SAC', 'E001-185', '2026-07-31', 6440.69, 1159.32), m=maestros(padron={PROV_MAT}))
    assert not fila(s2, 'E001-185').d.retencion


def test_retencion_desde_xml_del_comprobante():
    cr = leer_xml_retencion(xml_retencion('E001-989', '2026-07-31', PROV_MAT, 'MATERIALES SAC', [('E001-185', '2026-07-31', 7600.01, 228.00)]))
    assert cr.serie == 'E001' and cr.numero == '989' and cr.ruc_proveedor == PROV_MAT and cr.referencias[0].retencion == 228.0
    s, _ = sesion(xml_factura(PROV_MAT, 'MATERIALES SAC', 'E001-185', '2026-07-31', 6440.69, 1159.32), m=maestros(retenciones=[cr]))
    d = fila(s, 'E001-185').d
    assert d.retencion and (d.retencion_serie, d.retencion_numero, d.retencion_fecha) == ('E001', '989', '2026-07-31')
    assert d.glosa.endswith('- RETENCION')


def test_padron_txt():
    assert leer_padron([('p.txt', b'20100047218|BANCO|X\n20999999901|OTRO|Y\nRUC|NOMBRE\n')]) == {'20100047218', '20999999901'}


# ----------------------------------------------------------------------------- bloques
def test_bloques_ciclo_completo_y_excel():
    s, m = sesion(
        xml_factura(PROV_MAT, 'MATERIALES SAC', 'E001-185', '2026-07-31', 6440.69, 1159.32),
        xml_factura(PROV_CAJA, 'FERRE SAC', 'F010-60', '2026-07-23', 100),
        xml_factura(PROV_CAJA, 'FERRE SAC', 'F010-59', '2026-07-24', 100),
    )
    f1 = fila(s, 'E001-185')
    s.editar(f1.id, m, comision=4.30, retencion_serie='E001', retencion_numero='989', retencion_fecha='2026-07-31')
    ids = [f.id for f in s.lista('pendiente')]
    b = s.crear_bloque(ids, 'Prueba', m)
    assert b.codigo == 'B01' and not s.lista('pendiente')
    # orden: del más reciente al más antiguo; las de ajuste al final
    assert [s.filas[i].c.serie_numero for i in b.ids] == ['F010-59', 'F010-60', 'E001-185']
    wb = openpyxl.load_workbook(io.BytesIO(excel_bloque(s, b)))
    filas = list(wb.active.iter_rows(values_only=True))
    assert len(filas) == 3 and all(len(r) == 50 for r in filas)
    ultima = filas[-1]
    assert ultima[31] == '6031025' and ultima[33] == '104103' and ultima[30] == 'CON'    # AF, AH, AE
    assert ultima[22] == 3.40 and ultima[36] in (None, '')                              # W lleno, AK vacío
    assert ultima[44].endswith('- RETENCION + G.BANCARIO')
    rev = openpyxl.load_workbook(io.BytesIO(excel_revision_bloque(s, b, m.plan)))
    aj = [r for r in rev['AJUSTES DE PAGO'].iter_rows(values_only=True) if r and r[2] == 'E001-185'][0]
    assert aj[7] == 4.30 and aj[8] == 228.00 and aj[9] == 7376.31                       # comisión, retención, 104 final
    excel_mes(s, m.plan)
    # devolver y volver a crear
    s.devolver([f1.id])
    assert f1.d.estado == 'pendiente' and f1.id not in b.ids
    # confirmar con el Diario
    mem = construir_memoria(leer_diario(diario_excel([
        {'mes': '07', 'n': 5, 'ruc': PROV_CAJA, 'doc': '01', 'numero': '00F010-0000000000060', 'fecha': date(2026, 7, 23),
         'lineas': [('6560094', 100, 0), ('401111', 18, 0), ('1011', 0, 118)]}])))
    r = s.confirmar_con_diario(mem)
    assert r['confirmados'] == 1 and len(r['faltan']) == 1
    assert fila(s, 'F010-60').d.estado == 'contabilizado' and fila(s, 'F010-60').d.asiento == '07-5'


def test_no_deja_crear_bloque_en_rojo():
    s, m = sesion(xml_factura(PROV_CAJA, 'FERRE SAC', 'F010-80', '2026-07-20', 100))
    f = fila(s, 'F010-80')
    s.editar(f.id, m, cuenta_gasto='9999999')
    with pytest.raises(ValueError):
        s.crear_bloque([f.id], 'x', m)


def test_sin_tipo_de_cambio_es_rojo():
    m = maestros()
    m.tc = HistorialTC([])
    s, _ = sesion(xml_factura(PROV_CAJA, 'FERRE SAC', 'F010-81', '2026-07-20', 100), m=m)
    assert fila(s, 'F010-81').nivel == 'rojo'


def test_6399_bloqueada():
    s, m = sesion(xml_factura(PROV_CAJA, 'FERRE SAC', 'F010-82', '2026-07-20', 100))
    f = fila(s, 'F010-82')
    s.editar(f.id, m, cuenta_gasto='6399094')
    assert f.nivel == 'rojo'


def test_detraccion_sin_constancia_excluida():
    s, _ = sesion(xml_factura('20999999906', 'SERVICIOS SAC', 'F001-9', '2026-07-10', 1000, detraccion=('037', 12, 141.6)))
    assert fila(s, 'F001-9').d.estado == 'excluido'


# ----------------------------------------------------------------------------- persistencia
def test_sesion_json_y_almacen(tmp_path):
    s, m = sesion(xml_factura(PROV_CAJA, 'FERRE SAC', 'F010-90', '2026-07-20', 100))
    f = fila(s, 'F010-90')
    s.editar(f.id, m, cuenta_gasto='6560095', nota='revisado con el papel')
    b = s.crear_bloque([f.id], 'B', m)
    alm = Almacen(ruta_sqlite=str(tmp_path / 'r.db'))
    alm.guardar('sesion:2026-07', s.a_dict())
    s2 = Sesion.desde_dict(json.loads(json.dumps(alm.leer('sesion:2026-07'))))
    s2.procesar(m)
    f2 = fila(s2, 'F010-90')
    assert f2.d.cuenta_gasto == '6560095' and f2.d.gasto_manual and f2.d.estado == 'exportado' and s2.bloques[0].codigo == b.codigo


def test_entidades():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append([None, 'R.U.C.', 'NOMBRE O RAZON SOCIAL'])
    ws.append([None, '20999999901', 'MATERIALES SAC'])
    buf = io.BytesIO()
    wb.save(buf)
    assert leer_entidades(buf.getvalue()) == {'20999999901': 'MATERIALES SAC'}
